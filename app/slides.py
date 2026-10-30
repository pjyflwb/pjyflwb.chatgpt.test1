"""将经过校验的大纲渲染为可编辑的 PowerPoint。"""
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pydantic import BaseModel, Field

class Slide(BaseModel):
    """限制单页文本长度，减少溢出版面的情况。"""
    title: str = Field(min_length=1, max_length=60)
    bullets: list[str] = Field(min_length=1, max_length=5)

class Outline(BaseModel):
    """模型输出与本地生成器之间的数据契约。"""
    slides: list[Slide] = Field(min_length=3, max_length=12)

def render_ppt(outline: Outline) -> bytes:
    """采用宽屏深色主题生成 PPT，返回内存中的文件内容。"""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for number, item in enumerate(outline.slides, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("101D32")
        for top, text, size, color in [
            (0.7, item.title, 30, "7BE0C3"),
            (2.0, "\n\n".join("• " + point[:160] for point in item.bullets), 22, "ECF2FA"),
            (6.9, f"AI STUDIO   /   {number:02}", 11, "91A2B9"),
        ]:
            frame = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.7), Inches(4.7 if top == 2 else 0.7)).text_frame
            frame.word_wrap = True
            frame.text = text
            for paragraph in frame.paragraphs:
                paragraph.font.size = Pt(size)
                paragraph.font.color.rgb = RGBColor.from_string(color)
    output = BytesIO()
    prs.save(output)
    return output.getvalue()
