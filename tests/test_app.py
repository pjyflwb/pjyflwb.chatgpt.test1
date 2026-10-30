"""验证核心接口，避免依赖真实密钥或产生调用费用。"""
import unittest
from io import BytesIO
from unittest.mock import patch, AsyncMock
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from app.main import app

class AppTests(unittest.TestCase):
    """覆盖演示结果、文件验证及模型适配。"""
    def setUp(self):
        self.key = patch('app.config.API_KEY', '')
        self.key.start()
        self.addCleanup(self.key.stop)
        self.client = TestClient(app)

    def test_home_and_validation(self):
        self.assertEqual(self.client.get('/').status_code, 200)
        self.assertTrue(self.client.get('/api/status').json()['demo'])
        self.assertEqual(self.client.post('/api/images', json={'prompt':'  '}).status_code, 422)

    def test_demo_image(self):
        data = self.client.post('/api/images', json={'prompt':'森林'}).json()
        self.assertTrue(data['demo'])
        self.assertTrue(data['url'].startswith('data:image/png;base64,'))

    def test_upload_validation(self):
        bad = self.client.post('/api/vision', files={'file':('fake.png', b'bad', 'image/png')})
        self.assertEqual(bad.status_code, 422)
        stream = BytesIO()
        Image.new('RGB', (20,20)).save(stream, format='PNG')
        good = self.client.post('/api/vision', files={'file':('ok.png',stream.getvalue(),'image/png')})
        self.assertTrue(good.json()['demo'])

    def test_ppt_is_editable(self):
        response = self.client.post('/api/presentations', json={'prompt':'人工智能','pages':5})
        self.assertEqual(response.status_code, 200)
        ppt = Presentation(BytesIO(response.content))
        self.assertEqual(len(ppt.slides), 5)
        self.assertTrue(any('人工智能' in shape.text for shape in ppt.slides[0].shapes if shape.has_text_frame))

    def test_provider_image_mapping(self):
        with patch('app.config.API_KEY','test'), patch('app.main.request_ai',new=AsyncMock(return_value={'data':[{'b64_json':'AAAA'}]})):
            response = self.client.post('/api/images', json={'prompt':'森林'})
            self.assertFalse(response.json()['demo'])

    def test_invalid_outline(self):
        with patch('app.config.API_KEY','test'), patch('app.main.chat',new=AsyncMock(return_value='not json')):
            response = self.client.post('/api/presentations', json={'prompt':'主题'})
            self.assertEqual(response.status_code, 502)

if __name__ == '__main__':
    unittest.main()
